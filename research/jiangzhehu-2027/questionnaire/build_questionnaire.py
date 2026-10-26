# -*- coding: utf-8 -*-
"""Build the requirements-confirmation questionnaire DOCX."""

import os
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "微软雅黑"

BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEADER_FILL = "1F3864"
ALT_FILL = "F2F5FA"
BORDER_COLOR = "D9D9D9"

BOX = "\u25a1"

QUESTION_COL = Cm(12.0)
ANSWER_COL = Cm(5.0)

TBLPR_ORDER = [
    "w:tblStyle",
    "w:tblpPr",
    "w:tblOverlap",
    "w:bidiVisual",
    "w:tblStyleRowBandSize",
    "w:tblStyleColBandSize",
    "w:tblW",
    "w:jc",
    "w:tblCellSpacing",
    "w:tblInd",
    "w:tblBorders",
    "w:shd",
    "w:tblLayout",
    "w:tblCellMar",
    "w:tblLook",
    "w:tblCaption",
    "w:tblDescription",
]

TRPR_ORDER = [
    "w:cnfStyle",
    "w:divId",
    "w:gridBefore",
    "w:gridAfter",
    "w:wBefore",
    "w:wAfter",
    "w:cantSplit",
    "w:trHeight",
    "w:tblHeader",
    "w:tblCellSpacing",
    "w:jc",
    "w:hidden",
]

TCPR_ORDER = [
    "w:cnfStyle",
    "w:tcW",
    "w:gridSpan",
    "w:hMerge",
    "w:vMerge",
    "w:tcBorders",
    "w:shd",
    "w:noWrap",
    "w:tcMar",
    "w:textDirection",
    "w:tcFitText",
    "w:vAlign",
    "w:hideMark",
]


def reorder(element, order):
    """Reorder child elements to match the OOXML schema sequence."""
    index = {qn(name): pos for pos, name in enumerate(order)}
    children = list(element)
    children.sort(key=lambda child: index.get(child.tag, len(order)))
    for child in children:
        element.append(child)


def style_run(run, size=10.5, bold=False, color=BLACK, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), font)
    return run


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)
    reorder(tcPr, TCPR_ORDER)


def set_table_borders(table, color=BORDER_COLOR, size=4):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for tag in ("w:top", "w:left", "w:bottom", "w:right", "w:insideH", "w:insideV"):
        el = OxmlElement(tag)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)
    reorder(tblPr, TBLPR_ORDER)


def set_cell_margins(table, top=110, start=145, bottom=110, end=145):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblCellMar")):
        tblPr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for tag, value in (
        ("w:top", top),
        ("w:start", start),
        ("w:bottom", bottom),
        ("w:end", end),
    ):
        el = OxmlElement(tag)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)
    reorder(tblPr, TBLPR_ORDER)


def set_fixed_layout(table):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(old)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    reorder(tblPr, TBLPR_ORDER)


def set_row_height(row, twips, header=False):
    trPr = row._tr.get_or_add_trPr()
    split = OxmlElement("w:cantSplit")
    trPr.append(split)
    if header:
        head = OxmlElement("w:tblHeader")
        trPr.append(head)
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), str(twips))
    height.set(qn("w:hRule"), "atLeast")
    trPr.append(height)
    reorder(trPr, TRPR_ORDER)


def set_col_widths(table, widths):
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def add_field(paragraph, instruction, placeholder="1"):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " %s " % instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, text, end):
        run._r.append(el)
    return run


def build_document():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size in (("Title", 19), ("Heading 1", 13)):
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rFonts.set(qn(attr), FONT)
        pPr = style.element.get_or_add_pPr()
        for border in pPr.findall(qn("w:pBdr")):
            pPr.remove(border)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.footer_distance = Cm(1.2)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    for border in title._p.get_or_add_pPr().findall(qn("w:pBdr")):
        title._p.get_or_add_pPr().remove(border)
    style_run(title.add_run("江浙沪电气自动化应届生就业与生活成本调研问卷"), size=18, bold=True)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(6)
    intro.paragraph_format.line_spacing = 1.35
    style_run(
        intro.add_run(
            "这份问卷用于在正式调研开始前确定统计口径。你的回答会决定报告覆盖哪些城市、"
            "按什么口径统计薪资与居住成本，以及最终结论怎么判断。"
        ),
        size=10.5,
    )

    howto = doc.add_paragraph()
    howto.paragraph_format.space_after = Pt(16)
    howto.paragraph_format.line_spacing = 1.35
    style_run(
        howto.add_run(
            "填写方式：在方框里打勾即可，需要说明或补充的写在右侧空白栏。"
            "没有想法的题目可以留空，拿不准的写不确定。整份问卷大约需要十五分钟。"
        ),
        size=10.5,
    )

    sections = [
        (
            "第一部分 调研范围",
            [
                (
                    "目标城市的覆盖范围",
                    None,
                    [
                        "核心城市深度版：上海、苏州、南京、无锡、常州、南通、杭州、宁波、嘉兴、温州，其余城市出简版",
                        "全覆盖版：江苏十三市、浙江十一市加上海，单座城市的深度会相应下降",
                        "只做我指定的城市，请在右栏列出",
                        "其他分级方式，请在右栏说明",
                    ],
                ),
                (
                    "江苏部分的覆盖范围",
                    "江苏十三市：南京、无锡、徐州、常州、苏州、南通、连云港、淮安、盐城、扬州、镇江、泰州、宿迁",
                    [
                        "全部覆盖",
                        "去掉苏北五市（徐州、连云港、淮安、盐城、宿迁）",
                        "只保留苏南与沿江城市",
                        "其他，请在右栏说明",
                    ],
                ),
                (
                    "浙江部分的覆盖范围",
                    "浙江十一市：杭州、宁波、温州、嘉兴、湖州、绍兴、金华、衢州、舟山、台州、丽水",
                    [
                        "全部覆盖",
                        "去掉衢州、舟山、丽水",
                        "只保留环杭州湾五市（杭州、宁波、嘉兴、绍兴、湖州）",
                        "其他，请在右栏说明",
                    ],
                ),
                (
                    "上海是否细分到区",
                    "浦东张江、闵行、嘉定、松江、宝山、青浦、临港等区域的房租差异明显",
                    [
                        "需要细分到区",
                        "只按全市平均",
                        "其他，请在右栏说明",
                    ],
                ),
                (
                    "房租是否细分到区县一级",
                    None,
                    [
                        "需要，同一城市不同区差异可达一倍",
                        "只到城市级",
                        "其他，请在右栏说明",
                    ],
                ),
            ],
        ),
        (
            "第二部分 学历与就业方向",
            [
                (
                    "你的专业方向",
                    None,
                    [
                        "电气工程及其自动化",
                        "自动化",
                        "电气工程",
                        "其他，请在右栏填写",
                    ],
                ),
                (
                    "你当前所处的阶段",
                    "这一项决定报告重点看哪一届的数据",
                    [
                        "在校生，正在参加 2027 届秋招",
                        "在校生，离毕业还有一年以上",
                        "已毕业，正在找工作",
                        "已入职，想看自己的薪资是否偏低",
                        "其他，请在右栏说明",
                    ],
                ),
                (
                    "主要参考的学历档（可多选，请标出主档）",
                    None,
                    ["专科", "本科", "硕士", "博士"],
                ),
                (
                    "是否区分学校层次",
                    None,
                    [
                        "需要，区分 985 与 211、电力行业院校、普通本科",
                        "不需要，只按学历分档",
                        "其他，请在右栏说明",
                    ],
                ),
                (
                    "需要重点调研的行业方向（可多选）",
                    None,
                    [
                        "国家电网与南方电网（供电公司、电科院）",
                        "发电集团（华能、大唐、华电、国家电投、三峡、中广核）",
                        "电力装备（南瑞、许继、平高、思源、正泰、良信、汇川）",
                        "新能源（宁德时代、阳光电源、隆基、晶科、天合、远景、金风）",
                        "轨道交通（中车、地铁公司、中铁建）",
                        "外企自动化（西门子、ABB、施耐德、罗克韦尔、霍尼韦尔）",
                        "半导体与电子制造",
                        "汽车电子与新能源汽车",
                        "建筑电气、设计院与施工",
                        "公务员与事业单位",
                        "其他，请在右栏补充",
                    ],
                ),
                (
                    "目标岗位类型（可多选）",
                    None,
                    [
                        "研发与设计",
                        "现场调试与工程实施",
                        "生产工艺与制造",
                        "销售与技术支持",
                        "项目管理",
                        "运行维护",
                        "其他，请在右栏补充",
                    ],
                ),
                (
                    "是否区分用工性质",
                    None,
                    [
                        "需要，区分正式编制、劳务派遣、外包",
                        "不需要",
                        "其他，请在右栏说明",
                    ],
                ),
            ],
        ),
        (
            "第三部分 薪资统计口径",
            [
                (
                    "薪资以哪个口径为主（单选）",
                    None,
                    [
                        "税前月薪",
                        "税后到手月薪",
                        "全年总包，含年终奖",
                        "第一年实际到手总额",
                        "其他，请在右栏说明",
                    ],
                ),
                (
                    "是否需要拆解收入结构",
                    None,
                    [
                        "需要，拆成基本工资、绩效、年终奖、各类补贴、公积金",
                        "不需要",
                        "其他，请在右栏说明",
                    ],
                ),
                (
                    "统计指标偏好",
                    "平均值容易被少数高薪拉高，中位数与分位数更接近真实分布",
                    [
                        "中位数与分位数，即 P25、P50、P75",
                        "平均值",
                        "两者都要",
                        "其他，请在右栏说明",
                    ],
                ),
                (
                    "数据对应的时间口径（可多选）",
                    None,
                    [
                        "2026 届已入职的实际薪资",
                        "2027 届秋招进行中的 offer",
                        "近三年即 2024 至 2026 年的变化趋势",
                        "都要",
                    ],
                ),
                (
                    "是否需要税前税后换算",
                    "按各城市社保与公积金规则计算到手金额",
                    ["需要", "不需要", "其他，请在右栏说明"],
                ),
                (
                    "是否单列公积金差异",
                    "缴存比例与基数会明显影响实际收入与购房能力",
                    ["需要", "不需要", "其他，请在右栏说明"],
                ),
            ],
        ),
        (
            "第四部分 居住与日常开支",
            [
                (
                    "需要统计的租房类型（可多选）",
                    None,
                    [
                        "合租次卧",
                        "合租主卧",
                        "整租一居",
                        "整租两居",
                        "品牌公寓，如自如、泊寓",
                        "老小区或城中村",
                        "人才公寓",
                        "其他，请在右栏补充",
                    ],
                ),
                (
                    "租房位置的统计口径（可多选）",
                    None,
                    [
                        "市中心步行可达",
                        "地铁沿线通勤四十分钟内",
                        "产业园区附近",
                        "远郊通勤一小时以上",
                        "以上都要对比",
                    ],
                ),
                (
                    "租金之外需要计入的费用（可多选）",
                    None,
                    [
                        "押付方式与现金流占用，如押一付三",
                        "中介费",
                        "品牌公寓服务费",
                        "水电燃气与宽带",
                        "物业费",
                        "通勤交通费",
                        "全部计入",
                    ],
                ),
                (
                    "餐饮的统计口径（可多选）",
                    None,
                    [
                        "外卖实付均价",
                        "快餐堂食，如面馆、盖浇饭",
                        "公司食堂",
                        "自己做饭的食材成本",
                        "超市与菜市场价格",
                        "全部计入",
                    ],
                ),
                (
                    "其他生活开支是否需要覆盖（可多选）",
                    None,
                    [
                        "通勤交通",
                        "手机与网络",
                        "衣物与日用",
                        "社交与娱乐",
                        "医疗",
                        "商业保险",
                        "还贷或助学贷款",
                        "赡养父母",
                        "其他，请在右栏补充",
                    ],
                ),
                (
                    "是否需要计算每月净结余",
                    "净结余指月到手收入减去房租与生活开支之后剩下的钱",
                    [
                        "需要",
                        "不需要，只提供原始价格",
                        "其他，请在右栏说明",
                    ],
                ),
            ],
        ),
        (
            "第五部分 个人情况",
            [
                (
                    "住房情况",
                    None,
                    [
                        "江浙沪本地有房",
                        "本地有亲属可住",
                        "完全需要自己租房",
                        "其他，请在右栏说明",
                    ],
                ),
                (
                    "是否接受合租",
                    None,
                    ["接受", "不接受，必须整租", "视价格而定"],
                ),
                (
                    "是否有人可以分摊房租",
                    None,
                    ["有，两人分摊", "没有，一个人承担", "其他，请在右栏说明"],
                ),
                (
                    "是否接受驻场或长期出差",
                    "这类岗位补贴较高，但生活质量受影响",
                    ["接受", "不接受", "视补贴而定"],
                ),
                (
                    "加班强度的底线",
                    None,
                    [
                        "接受高强度加班",
                        "只接受适度加班",
                        "必须双休",
                        "其他，请在右栏说明",
                    ],
                ),
                (
                    "是否把考公、考编、考研作为并行选项一并调研",
                    None,
                    ["是", "否"],
                ),
                (
                    "购房计划的时间",
                    None,
                    ["三年内", "五年内", "不确定", "不打算"],
                ),
            ],
        ),
        (
            "第六部分 数据来源与交付方式",
            [
                (
                    "数据来源偏好",
                    None,
                    [
                        "只用官方与平台公开报告，来源硬但覆盖面窄",
                        "官方报告、平台数据与个体样本混合，并标注证据等级",
                        "能参考即可，不特别要求",
                    ],
                ),
                (
                    "是否接受个体样本数据",
                    "包括招聘爆料与社交平台的一手分享，这类数据只能作为参考",
                    [
                        "接受，但必须标注为参考",
                        "不接受",
                        "其他，请在右栏说明",
                    ],
                ),
                (
                    "是否同意直接运行 GitHub 上的采集工具",
                    "例如 BOSS 直聘职位薪资抓取、链家与豆瓣租房抓取，可能受平台反爬限制",
                    [
                        "同意",
                        "不同意，只用公开报告",
                        "先做可行性验证再决定",
                    ],
                ),
                (
                    "期望的交付形式（可多选）",
                    None,
                    [
                        "Word 报告",
                        "Excel 表格，可自行筛选排序",
                        "可调参数的可视化页面",
                        "Markdown 文档",
                        "全部都要",
                    ],
                ),
                (
                    "报告篇幅",
                    None,
                    [
                        "精简，十页以内，结论为主",
                        "详细，含方法与数据来源",
                        "越详细越好",
                    ],
                ),
                (
                    "是否需要给出推荐城市排序与理由",
                    None,
                    ["需要", "不需要，只提供数据"],
                ),
                (
                    "其他需要补充的问题或特别关心的事项",
                    None,
                    ["!请在右栏填写"],
                ),
            ],
        ),
    ]

    question_number = 0

    for heading, items in sections:
        head = doc.add_paragraph(style="Heading 1")
        head.paragraph_format.space_before = Pt(18)
        head.paragraph_format.space_after = Pt(8)
        head.paragraph_format.keep_with_next = True
        style_run(head.add_run(heading), size=13, bold=True)

        table = doc.add_table(rows=1 + len(items), cols=2)
        table.autofit = False
        set_fixed_layout(table)
        set_table_borders(table)
        set_cell_margins(table)

        header_row = table.rows[0]
        set_row_height(header_row, 560, header=True)
        for cell, text in zip(header_row.cells, ["问题与选项", "你的回答"]):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            shade_cell(cell, HEADER_FILL)
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(0)
            style_run(para.add_run(text), size=10.5, bold=True, color=WHITE)

        for index, (question, note, options) in enumerate(items):
            question_number += 1
            row = table.rows[index + 1]
            set_row_height(row, 900)

            q_cell, a_cell = row.cells
            q_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            a_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            fill = ALT_FILL if index % 2 else None
            if fill:
                shade_cell(q_cell, fill)
                shade_cell(a_cell, fill)

            first = q_cell.paragraphs[0]
            first.paragraph_format.space_after = Pt(3)
            first.paragraph_format.line_spacing = 1.25
            style_run(
                first.add_run("%d. %s" % (question_number, question)),
                size=10.5,
                bold=True,
            )

            if note:
                note_para = q_cell.add_paragraph()
                note_para.paragraph_format.space_after = Pt(3)
                note_para.paragraph_format.line_spacing = 1.2
                style_run(note_para.add_run("说明：%s" % note), size=9)

            for option in options:
                para = q_cell.add_paragraph()
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.2
                para.paragraph_format.left_indent = Cm(0.35)
                if option.startswith("!"):
                    text = option[1:]
                else:
                    text = "%s %s" % (BOX, option)
                style_run(para.add_run(text), size=10)

            blank = a_cell.paragraphs[0]
            blank.paragraph_format.space_after = Pt(0)
            style_run(blank.add_run(""), size=10.5)

    closing = doc.add_paragraph()
    closing.paragraph_format.space_before = Pt(16)
    closing.paragraph_format.line_spacing = 1.35
    style_run(
        closing.add_run("填写完成后，把这份文件发回即可，我会按你的选择开始采集。"),
        size=10.5,
    )

    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(footer_para.add_run("第 "), size=9)
    style_run(add_field(footer_para, "PAGE"), size=9)
    style_run(footer_para.add_run(" 页，共 "), size=9)
    style_run(add_field(footer_para, "NUMPAGES"), size=9)
    style_run(footer_para.add_run(" 页"), size=9)

    props = doc.core_properties
    props.title = "江浙沪电气自动化应届生就业与生活成本调研 需求确认问卷"
    props.author = "Codex"
    props.last_modified_by = "Codex"

    return doc


def main():
    out_path = sys.argv[1] if len(sys.argv) > 1 else "questionnaire.docx"
    out_dir = os.path.dirname(os.path.abspath(out_path))
    if out_dir and not os.path.isdir(out_dir):
        os.makedirs(out_dir, exist_ok=True)
    doc = build_document()
    doc.save(out_path)
    print("saved: %s" % out_path)


if __name__ == "__main__":
    main()
